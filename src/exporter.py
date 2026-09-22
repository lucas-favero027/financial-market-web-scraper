"""Persistence functions for raw snapshots, CSV files, and Excel output."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Font, PatternFill

PROJECT_ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = PROJECT_ROOT / "data" / "raw"
PROCESSED_DIR = PROJECT_ROOT / "data" / "processed"


def save_raw_snapshot(
    payloads: dict[str, Any],
    errors: dict[str, str],
    collected_at: datetime,
) -> Path:
    """Save source responses together for traceability."""
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = collected_at.strftime("%Y%m%d_%H%M%S")
    path = RAW_DIR / f"market_snapshot_{timestamp}.json"
    content = {
        "collected_at": collected_at.isoformat(),
        "sources": payloads,
        "errors": errors,
    }
    path.write_text(json.dumps(content, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def _write_csv(data: pd.DataFrame, path: Path) -> None:
    """Write a UTF-8 CSV with ISO dates."""
    data.to_csv(
        path,
        index=False,
        encoding="utf-8-sig",
        date_format="%Y-%m-%d %H:%M:%S",
    )


def _format_worksheet(worksheet) -> None:  # type: ignore[no-untyped-def]
    """Apply a compact, readable format to one generated worksheet."""
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_cells in worksheet.columns:
        letter = column_cells[0].column_letter
        max_length = max(len(str(cell.value or "")) for cell in column_cells[:200])
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 32)

    headers = {cell.value: cell.column for cell in worksheet[1]}
    for row in range(2, worksheet.max_row + 1):
        for field in ("price", "open", "high", "low"):
            if field in headers:
                worksheet.cell(row, headers[field]).number_format = 'R$ #,##0.00'
        for field in ("change_percent", "composition_percent", "value"):
            if field in headers:
                worksheet.cell(row, headers[field]).number_format = "0.000"
        for field in ("volume", "theoretical_quantity"):
            if field in headers:
                worksheet.cell(row, headers[field]).number_format = "#,##0.########"
        if "updated_at" in headers:
            worksheet.cell(row, headers["updated_at"]).number_format = (
                "dd/mm/yyyy hh:mm:ss"
            )


def export_market_data(
    stocks: pd.DataFrame,
    fiis: pd.DataFrame,
    crypto: pd.DataFrame,
    indicators: pd.DataFrame,
    market_data: pd.DataFrame,
) -> dict[str, Path]:
    """Export category CSVs, a consolidated CSV, and one Excel workbook."""
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    datasets = {
        "stocks": stocks,
        "fiis": fiis,
        "crypto": crypto,
        "economic_indicators": indicators,
        "market_data": market_data,
    }
    paths: dict[str, Path] = {}
    for name, data in datasets.items():
        path = PROCESSED_DIR / f"{name}.csv"
        _write_csv(data, path)
        paths[name] = path

    excel_path = PROCESSED_DIR / "market_data.xlsx"
    sheet_data = {
        "Market Data": market_data,
        "Stocks": stocks,
        "FIIs": fiis,
        "Crypto": crypto,
        "Indicators": indicators,
    }
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        for sheet_name, data in sheet_data.items():
            data.to_excel(writer, index=False, sheet_name=sheet_name)
            _format_worksheet(writer.sheets[sheet_name])
    paths["excel"] = excel_path
    return paths
